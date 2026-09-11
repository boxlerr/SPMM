"""Arma la planilla de trabas para que Lucas la conteste, con el estado de HOY.

Es la segunda vuelta de lo del 1/9: en vez de corregir traba por traba adentro de
cada planificación, se juntan todas las que quedan, se mandan en un Excel y se
cierran de una. La primera planilla se armó a mano y se perdió el script; esta
sale de la base, así que se puede volver a correr cuando quede otra tanda.

QUÉ PREGUNTA (sale de mirar el taller, no de una lista escrita a mano)

  1. EN QUÉ MÁQUINA SE HACE — procesos que el planificador da por trabajo de
     máquina pero para los que no sabe cuál: no la tienen cargada en Recursos y el
     nombre tampoco la delata. Terminan planificados sin reservar máquina, que es
     el aviso «no se sabe en qué máquina se hace».

  2. LA MÁQUINA NO ACEPTA AL QUE LO HACE — el proceso pide una categoría que
     ninguna de sus máquinas admite. Es la traba que Lucas leyó al revés en la
     reunión del 10/9: la pantalla dice «la máquina no acepta el rango que pide» y
     él entendía que el problema era el proceso.

  3. NADIE PUEDE HACERLO — procesos sin ninguna categoría cargada. No es que falte
     gente: falta el dato, así que no hay nadie a quien dárselo.

Solo mira procesos EN USO en OT abiertas: preguntar por el resto es hacerle perder
el tiempo. Las respuestas se cargan después con un script como
`cargar_planilla_trabas_20260902.py`, que deja asentado de qué respuesta sale cada
cambio.

    venv/bin/python -m backend.scripts.planilla_trabas
    venv/bin/python -m backend.scripts.planilla_trabas --salida /ruta/archivo.xlsx
"""
import asyncio
import os
import re
import sys
from datetime import date

import asyncpg
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from backend.application.PlanificacionService import (
    familia_from_maquina,
    familia_requerida_from_proceso,
    proceso_usa_maquina,
)

SALIDA = None
if "--salida" in sys.argv:
    SALIDA = sys.argv[sys.argv.index("--salida") + 1]
SALIDA = SALIDA or os.path.join(
    "tmp", f"Trabas del planificador - Metlo - {date.today():%Y-%m-%d}.xlsx")

# Paleta de la planilla que Lucas ya usó. No es decoración: la contestó una vez con
# estos colores (amarillo = te toca a vos, verde = ya está) y cambiarlos ahora es
# volver a explicarle la planilla.
AZUL, VERDE, AMARILLO, ROSA, VERDE_CLARO = "FF2E5C8A", "FF217346", "FFFFF2CC", "FFFFC7CE", "FFE2EFDA"
AZUL_OSCURO = "FF1F3864"
BORDE = Border(*[Side(style="thin", color="FFD0D0D0")] * 4)


def _veces(n):
    return "1 vez" if n == 1 else f"{n} veces"


def _ots(n):
    return "1 OT abierta" if n == 1 else f"{n} OT abiertas"


def _url():
    load_dotenv()
    u = re.sub(r"^postgresql\+\w+://", "postgresql://", os.getenv("SUPABASE_DB_URL")).split("?")[0]
    return u.replace(":5432/", ":6543/")


# ---------------------------------------------------------------------------
# Los datos
# ---------------------------------------------------------------------------
SQL_USO = """
with abiertas as (
    select id from orden_trabajo
    where coalesce(finalizadototal,0) = 0
      and (fecha_entrega is null or extract(year from fecha_entrega) <= 1950)
),
usados as (
    select p.id_proceso, count(*) as veces, count(distinct p.id_orden_trabajo) as ots
    from orden_trabajo_proceso p join abiertas a on a.id = p.id_orden_trabajo
    group by 1
)
select pr.id, pr.nombre, u.veces, u.ots,
       (select string_agg(distinct r.nombre, ', ' order by r.nombre)
          from rango_proceso rp join rango r on r.id = rp.id_rango
         where rp.id_proceso = pr.id) as rangos,
       (select string_agg(distinct m.nombre, ', ' order by m.nombre)
          from proceso_maquinaria pm join maquinaria m on m.id = pm.id_maquinaria
         where pm.id_proceso = pr.id) as maquinas,
       (select string_agg(distinct r2.nombre, ', ' order by r2.nombre)
          from proceso_maquinaria pm2
          join rango_maquinaria rm on rm.id_maquinaria = pm2.id_maquinaria
          join rango r2 on r2.id = rm.id_rango
         where pm2.id_proceso = pr.id) as maquinas_aceptan,
       (select string_agg(distinct ot.id_otvieja::text, ', ')
          from orden_trabajo_proceso p2
          join orden_trabajo ot on ot.id = p2.id_orden_trabajo
         where p2.id_proceso = pr.id and p2.id_orden_trabajo in (select id from abiertas)) as ot_afectadas
from usados u join proceso pr on pr.id = u.id_proceso
order by u.veces desc, pr.nombre
"""


# Las preguntas que NO son de procesos. Van en la misma planilla a pedido de Julián
# (11/09): «toda pregunta que tengamos tiene que responderla él acá».
SQL_OT_VIEJAS = """
select ot.id_otvieja, c.nombre as cliente,
       coalesce(a.descripcion, ot.observaciones) as que_es,
       ot.fecha_entrada::date as entrada, ot.unidades, ot.cantidad_entregada
  from orden_trabajo ot
  left join cliente c on c.id = ot.id_cliente
  left join articulo a on a.id = ot.id_articulo
 where coalesce(ot.finalizadototal, 0) = 0 and ot.fecha_entrada < '2026-01-01'
 order by ot.id_otvieja
"""

# Máquinas que hoy puede usar UNA persona o NINGUNA.
#
# No alcanza con «tiene una sola categoría»: de las 27 máquinas, 19 la tienen y casi
# todas están perfectas —la categoría es OFICIAL y hay seis oficiales—. Preguntar por
# esas 19 es hacerle perder el tiempo al taller con dieciséis preguntas que no
# cambian nada.
#
# Lo que importa es cuánta GENTE REAL queda del otro lado del cruce. Con ese criterio
# quedan las que de verdad se pueden frenar, e incluye el peor caso, que el criterio
# anterior dejaba afuera: los tres TORNO CNC piden OFICIAL ESPECIALIZADO o TÉCNICO y
# esos rangos no los tiene ninguna persona, así que están en cero.
SQL_MAQUINAS_TRABADAS = """
select m.nombre,
       string_agg(distinct r.nombre, ' o ' order by r.nombre) as rangos,
       count(distinct o.id) as gente
  from maquinaria m
  join rango_maquinaria rm on rm.id_maquinaria = m.id
  join rango r on r.id = rm.id_rango
  left join operario_rango orr on orr.id_rango = r.id
  left join operario o on o.id = orr.id_operario and o.disponible
 group by m.id, m.nombre
having count(distinct o.id) <= 1
 order by count(distinct o.id), m.nombre
"""

# Los OFICIAL y qué familia de máquina toca cada uno. Es el material de la pregunta
# que quedó del 28/08: que un oficial soldador no sea lo mismo que un oficial tornero.
SQL_OFICIALES = """
select o.id, o.nombre || ' ' || coalesce(o.apellido, '') as quien
  from operario o
  join operario_rango orr on orr.id_operario = o.id
  join rango r on r.id = orr.id_rango
 where o.disponible and r.nombre = 'OFICIAL'
 order by 2
"""

SQL_PROCESOS_DE = """
select distinct p.nombre
  from operario_rango orr
  join rango_proceso rp on rp.id_rango = orr.id_rango
  join proceso p on p.id = rp.id_proceso
 where orr.id_operario = $1
   and not exists (select 1 from operario_proceso_skill s
                    where s.id_operario = orr.id_operario
                      and s.id_proceso = p.id and not s.habilitado)
"""


async def _datos(c):
    filas = [dict(f) for f in await c.fetch(SQL_USO)]
    maquinas = [f["nombre"] for f in await c.fetch("select nombre from maquinaria order by nombre")]
    # Con el código y las categorías: el código es lo que mira `familia_from_maquina`.
    maquinas_detalle = [dict(f) for f in await c.fetch("""
        select m.nombre, m.cod_maquina as cod,
               coalesce(array_agg(r.nombre) filter (where r.nombre is not null), '{}') as rangos
          from maquinaria m
          left join rango_maquinaria rm on rm.id_maquinaria = m.id
          left join rango r on r.id = rm.id_rango
         group by m.id, m.nombre, m.cod_maquina
         order by m.nombre""")]
    rangos = [f["nombre"] for f in await c.fetch("select nombre from rango order by nombre")]

    # Las máquinas de cada familia, y qué categorías acepta cada una.
    #
    # ESTO ES LO QUE FALTABA, Y ERA EL AGUJERO MÁS GRANDE DE LA PLANILLA.
    #
    # Antes el cruce de categorías se calculaba mirando SÓLO `proceso_maquinaria`,
    # que tiene 13 filas en toda la base. Pero el planificador no usa esa tabla para
    # saber en qué máquina va un trabajo: deduce la FAMILIA DEL NOMBRE del proceso
    # («preparación de soldadora TIG» → familia SOLDADORA_TIG). Así que la planilla
    # sólo preguntaba por los poquísimos procesos que además tenían fila en el
    # catálogo, y se perdía el resto.
    #
    # Cuánto se perdía, medido el 11/9: preguntaba por 2 cruces de categoría y los
    # que existen son 5, que son 37 pasadas de OT abiertas. El más grande de todos
    # —PREPARACIÓN DE SOLDADORA TIG, 19 pasadas, una sola pregunta de sí o no— nunca
    # llegó a la planilla que el taller contestó.
    #
    # Ahora se resuelve igual que el planificador: primero el catálogo si lo hay,
    # y si no, las máquinas de la familia que dice el nombre.
    por_familia = {}
    for m in maquinas_detalle:
        fam = familia_from_maquina(m["nombre"], m["cod"] or "")
        if fam:
            por_familia.setdefault(fam, []).append(m)

    sin_maquina, rango_imposible, sin_rango = [], [], []
    for f in filas:
        tercerizado = any(x in (f["rangos"] or "") for x in ("TERCERIZADO", "EXTERNO"))
        if not f["rangos"]:
            sin_rango.append(f)
            continue
        if not proceso_usa_maquina(f["nombre"], tercerizado):
            continue

        familia = familia_requerida_from_proceso(f["nombre"])
        if f["maquinas"]:
            sus_maquinas = [{"nombre": n.strip()} for n in f["maquinas"].split(",") if n.strip()]
            acepta = {r.strip() for r in (f["maquinas_aceptan"] or "").split(",") if r.strip()}
        elif familia and por_familia.get(familia):
            sus_maquinas = por_familia[familia]
            acepta = {r for m in sus_maquinas for r in m["rangos"]}
        else:
            # Ni catálogo ni familia con máquinas: no se sabe dónde se hace.
            sin_maquina.append(f)
            continue

        pide = {r.strip() for r in (f["rangos"] or "").split(",") if r.strip()}
        if acepta and not (pide & acepta):
            # El texto de la pregunta tiene que nombrar las máquinas de verdad, que
            # ahora pueden venir de la familia y no del catálogo.
            f["maquinas"] = ", ".join(sorted(m["nombre"] for m in sus_maquinas))
            f["maquinas_aceptan"] = ", ".join(sorted(acepta))
            rango_imposible.append(f)
    return sin_maquina, rango_imposible, sin_rango, maquinas, rangos


# ---------------------------------------------------------------------------
# El Excel
# ---------------------------------------------------------------------------
def _titulo(ws, celda, texto, fondo=AZUL, tam=11, color="FFFFFFFF"):
    ws[celda] = texto
    ws[celda].fill = PatternFill("solid", fgColor=fondo)
    ws[celda].font = Font(bold=True, size=tam, color=color)


def _hoja_empeza(wb, total):
    ws = wb.create_sheet("EMPEZÁ ACÁ")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 95

    _titulo(ws, "B2", "Trabas del planificador — segunda tanda", AZUL_OSCURO, 16)
    ws["B3"] = f"Metalúrgica Longchamps · {date.today():%d de %B de %Y}"
    ws["B3"].font = Font(size=11, color="FF666666")

    _titulo(ws, "B5", "Qué hay que hacer", AZUL, 12)
    pasos = [
        ("1", "Andá a la hoja de abajo que dice PREGUNTAS."),
        ("2", "Contestá la columna amarilla. Hacé clic en la celda y elegí de la lista: "
              "no hace falta escribir nada."),
        ("3", "Si no sabés, elegí «No sé». Es una respuesta válida y me sirve igual."),
        ("4", "Cuando termines, mandámela. Con eso cargo todo de una y dejan de aparecer "
              "estas trabas al planificar."),
    ]
    for i, (n, txt) in enumerate(pasos, start=6):
        ws[f"B{i}"] = f"  {n}"
        ws[f"B{i}"].font = Font(bold=True, size=11)
        ws[f"C{i}"] = txt
        ws[f"C{i}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30

    ws["B11"] = f"Son {total} preguntas. Cada una se contesta con un clic."
    ws["B11"].font = Font(bold=True, size=12)

    _titulo(ws, "B13", "Qué te estoy preguntando", AZUL, 12)
    bloques = [
        ("EN QUÉ MÁQUINA SE HACE",
         "El sistema sabe que el trabajo usa máquina, pero no cuál. Los planifica igual "
         "y no reserva la máquina, así que dos trabajos pueden caer juntos en la misma."),
        ("LA MÁQUINA NO ACEPTA AL QUE LO HACE",
         "El trabajo pide una categoría que ninguna de sus máquinas admite. Es la traba "
         "que aparece como «sus máquinas no aceptan el rango que pide»."),
        ("NADIE PUEDE HACERLO",
         "El trabajo no tiene ninguna categoría cargada, así que no hay a quién dárselo. "
         "No falta gente: falta el dato."),
    ]
    fila = 14
    for nombre, expl in bloques:
        ws[f"B{fila}"] = nombre
        ws[f"B{fila}"].font = Font(bold=True, size=11)
        ws[f"C{fila}"] = expl
        ws[f"C{fila}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[fila].height = 32
        fila += 1

    ws[f"B{fila+1}"] = ("Al final hay una hoja «Detalle técnico» con el número de cada trabajo y "
                        "en qué OT aparece. No hace falta que la mires.")
    ws[f"B{fila+1}"].font = Font(italic=True, size=10, color="FF666666")
    return ws


def _hoja_listas(wb, opciones):
    """Una columna por lista desplegable. Van acá y no en la fórmula porque una lista
    escrita a mano adentro del `formula1` se rompe apenas una opción tiene una coma
    —y los nombres de máquina las tienen—."""
    ws = wb.create_sheet("Listas")
    ws.sheet_state = "hidden"
    for i, (clave, valores) in enumerate(opciones.items(), start=1):
        col = get_column_letter(i)
        ws[f"{col}1"] = clave
        for j, v in enumerate(valores, start=2):
            ws[f"{col}{j}"] = v
    return ws


def _hoja_preguntas(wb, grupos, opciones):
    ws = wb.create_sheet("PREGUNTAS")
    ws.sheet_view.showGridLines = False
    anchos = {"A": 5, "B": 34, "C": 64, "D": 34, "E": 58}
    for col, w in anchos.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    total = sum(len(g["filas"]) for g in grupos)
    _titulo(ws, "A1", "Contestá la columna amarilla eligiendo de la lista", AZUL_OSCURO, 16)
    ws["A2"] = f'=CONCATENATE("Contestadas: ",COUNTA(D5:D1000)," de {total}")'
    ws["A2"].fill = PatternFill("solid", fgColor=VERDE_CLARO)
    ws["A2"].font = Font(bold=True, size=12)

    for celda, txt, fondo in (("B3", "EL TRABAJO", AZUL), ("C3", "LA PREGUNTA", AZUL),
                              ("D3", "TU RESPUESTA", VERDE), ("E3", "¿Por qué te pregunto esto?", AZUL)):
        _titulo(ws, celda, txt, fondo, 11)

    fila, n = 4, 0
    validaciones = []
    for grupo in grupos:
        ws.merge_cells(f"A{fila}:E{fila}")
        _titulo(ws, f"A{fila}", f"   {grupo['titulo']}   ·   {grupo['bajada']}", AZUL, 12)
        ws.row_dimensions[fila].height = 22
        fila += 1
        desde = fila
        for f in grupo["filas"]:
            n += 1
            ws[f"A{fila}"] = n
            ws[f"A{fila}"].fill = PatternFill("solid", fgColor=ROSA)
            ws[f"A{fila}"].font = Font(bold=True, size=11)
            ws[f"A{fila}"].alignment = Alignment(horizontal="center")
            ws[f"B{fila}"] = f["trabajo"]
            ws[f"B{fila}"].font = Font(bold=True, size=11)
            ws[f"C{fila}"] = f["pregunta"]
            ws[f"D{fila}"].fill = PatternFill("solid", fgColor=AMARILLO)
            ws[f"D{fila}"].font = Font(bold=True, size=11)
            ws[f"E{fila}"] = f["porque"]
            ws[f"E{fila}"].font = Font(size=9, color="FF666666")
            for col in "BCDE":
                ws[f"{col}{fila}"].alignment = Alignment(wrap_text=True, vertical="top")
                ws[f"{col}{fila}"].border = BORDE
            ws.row_dimensions[fila].height = 46
            fila += 1
        validaciones.append((grupo["lista"], desde, fila - 1))

    for clave, desde, hasta in validaciones:
        col = get_column_letter(list(opciones).index(clave) + 1)
        dv = DataValidation(
            type="list",
            formula1=f"Listas!${col}$2:${col}${len(opciones[clave]) + 1}",
            allow_blank=True, showDropDown=False)
        dv.error = "Elegí una opción de la lista."
        ws.add_data_validation(dv)
        dv.add(f"D{desde}:D{hasta}")
    return ws


def _hoja_detalle(wb, grupos):
    ws = wb.create_sheet("Detalle técnico")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 60
    for i, t in enumerate(["#", "id proceso", "proceso", "pasadas",
                           "categorías que pide", "máquinas cargadas", "OT abiertas donde aparece"], start=1):
        c = ws.cell(row=1, column=i, value=t)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.font = Font(bold=True, color="FFFFFFFF")
    fila, n = 2, 0
    for grupo in grupos:
        for f in grupo["filas"]:
            n += 1
            d = f["dato"]
            ws.cell(row=fila, column=1, value=n)
            ws.cell(row=fila, column=2, value=d["id"])
            ws.cell(row=fila, column=3, value=d["nombre"])
            ws.cell(row=fila, column=4, value=d["veces"])
            ws.cell(row=fila, column=5, value=d["rangos"] or "—")
            ws.cell(row=fila, column=6, value=d["maquinas"] or "—")
            ws.cell(row=fila, column=7, value=(d["ot_afectadas"] or "")[:900])
            fila += 1
    return ws


async def main():
    c = await asyncpg.connect(_url(), statement_cache_size=0)
    try:
        sin_maquina, rango_imposible, sin_rango, maquinas, rangos = await _datos(c)
        ot_viejas = [dict(f) for f in await c.fetch(SQL_OT_VIEJAS)]
        reservadas = [dict(f) for f in await c.fetch(SQL_MAQUINAS_TRABADAS)]
        oficiales = [dict(f) for f in await c.fetch(SQL_OFICIALES)]
        familias_de = {}
        for o in oficiales:
            nombres = [f["nombre"] for f in await c.fetch(SQL_PROCESOS_DE, o["id"])]
            familias_de[o["id"]] = sorted(
                {familia_requerida_from_proceso(n) for n in nombres} - {""})
    finally:
        await c.close()

    opciones = {
        # «A mano» y «Se manda a hacer afuera» son respuestas de verdad y hay que
        # ofrecerlas: la planilla del 2/9 destapó cuatro trabajos que no usaban
        # ninguna máquina (pulido, engomado, decapado, pegado de goma) y que el
        # sistema venía buscándoles una.
        "maquinas": ["A mano, sin máquina", "Se manda a hacer afuera"] + maquinas + ["No sé"],
        "rangos": rangos + ["Se manda a hacer afuera", "No lo hace nadie", "No sé"],
        "arreglo": ["Cambiar la categoría que pide el trabajo",
                    "Habilitar esa categoría en la máquina",
                    "Se hace en otra máquina", "No sé"],
    }

    grupos = []
    if sin_maquina:
        grupos.append({
            "titulo": "EN QUÉ MÁQUINA SE HACE",
            "bajada": "El sistema sabe que usa máquina pero no cuál, así que no la reserva",
            "lista": "maquinas",
            "filas": [{
                "trabajo": f["nombre"].title(),
                "pregunta": f"«{f['nombre'].title()}»: ¿en qué máquina se hace?",
                "porque": (f"Aparece {_veces(f['veces'])} en {_ots(f['ots'])}. Hoy se planifica "
                           f"sin reservar máquina, así que puede caer junto con otro trabajo "
                           f"en la misma."),
                "dato": f,
            } for f in sin_maquina],
        })
    if rango_imposible:
        grupos.append({
            "titulo": "LA MÁQUINA NO ACEPTA AL QUE LO HACE",
            "bajada": "El trabajo pide una categoría que ninguna de sus máquinas admite",
            "lista": "arreglo",
            "filas": [{
                "trabajo": f["nombre"].title(),
                "pregunta": (f"«{f['nombre'].title()}» pide {f['rangos']}, pero "
                             f"{f['maquinas']} solo acepta {f['maquinas_aceptan']}. ¿Cómo lo arreglamos?"),
                "porque": (f"Aparece {_veces(f['veces'])} en {_ots(f['ots'])}. Con esto sin "
                           f"resolver, el plan no sale."),
                "dato": f,
            } for f in rango_imposible],
        })
    if sin_rango:
        grupos.append({
            "titulo": "NADIE PUEDE HACERLO",
            "bajada": "El trabajo no tiene ninguna categoría cargada",
            "lista": "rangos",
            "filas": [{
                "trabajo": f["nombre"].title(),
                "pregunta": f"«{f['nombre'].title()}»: ¿qué categoría lo hace?",
                "porque": (f"Aparece {_veces(f['veces'])} en {_ots(f['ots'])} y no tiene "
                           f"categoría, así que no hay a quién dárselo y la OT no se puede "
                           f"planificar."),
                "dato": f,
            } for f in sin_rango],
        })

    # ── Lo que no es un proceso pero también hay que preguntarle ──────────────
    if ot_viejas:
        opciones["ot"] = ["Ya se entregó: cerrarla", "Falta trabajo de verdad",
                          "Se cancela", "No sé"]
        grupos.append({
            "titulo": "LAS ÓRDENES VIEJAS QUE SIGUEN ABIERTAS",
            "bajada": "Son de 2025 y figuran abiertas; el planificador les sigue haciendo lugar",
            "lista": "ot",
            "filas": [{
                "trabajo": f"OT {f['id_otvieja']}",
                "pregunta": f"La OT {f['id_otvieja']} sigue abierta desde {f['entrada']:%d/%m/%Y}. ¿Qué hacemos?",
                "porque": (f"{(f['cliente'] or 'sin cliente')[:34]} · {(f['que_es'] or '')[:56]} · "
                           f"entregadas {f['cantidad_entregada'] or 0} de {f['unidades'] or 0}. "
                           f"Mientras siga abierta, el plan le reserva máquina y gente."),
                "dato": {"id": f["id_otvieja"], "nombre": f"OT {f['id_otvieja']}",
                         "veces": 0, "rangos": "", "maquinas": "", "ot_afectadas": ""},
            } for f in ot_viejas],
        })

    if reservadas:
        opciones["reservada"] = ["Sí: es a propósito, dejalo así",
                                 "No: también la puede usar un OFICIAL",
                                 "No: también la puede usar un MEDIO OFICIAL",
                                 "No: también la puede usar un OPERARIO CALIFICADO",
                                 "No sé"]
        grupos.append({
            "titulo": "MÁQUINAS QUE HOY PUEDE USAR UNA PERSONA O NINGUNA",
            "bajada": "Si esa persona falta, la máquina queda parada y el trabajo no entra",
            "lista": "reservada",
            "filas": [{
                "trabajo": f["nombre"],
                "pregunta": (f"A «{f['nombre']}» hoy la puede usar "
                             + ("NADIE" if f["gente"] == 0 else "una sola persona")
                             + f" (pide {f['rangos']}). ¿Es a propósito?"),
                "porque": ("Si es a propósito no se toca — pasa lo mismo con la plegadora. Si no, "
                           + ("esa máquina hoy no la puede agarrar nadie y el trabajo que la "
                              "necesita se planifica sin reservarla."
                              if f["gente"] == 0 else
                              "cuando esa persona falta o está ocupada la máquina queda parada.")),
                "dato": {"id": 0, "nombre": f["nombre"], "veces": f["gente"],
                         "rangos": f["rangos"], "maquinas": f["nombre"], "ot_afectadas": ""},
            } for f in reservadas],
        })

    if oficiales:
        opciones["especialidad"] = ["Hace de todo un poco", "Es sobre todo tornero",
                                    "Es sobre todo fresador", "Es sobre todo soldador",
                                    "Es sobre todo plegador", "Es sobre todo rectificador",
                                    "Otra cosa: lo hablamos", "No sé"]
        grupos.append({
            "titulo": "QUÉ HACE CADA OFICIAL",
            "bajada": "Hoy «OFICIAL» es una sola categoría: un soldador y un tornero valen igual",
            "lista": "especialidad",
            "filas": [{
                "trabajo": o["quien"].strip(),
                "pregunta": f"{o['quien'].strip()}: ¿en qué es especialista?",
                "porque": ("Hoy el sistema los trata a todos igual, así que le puede dar un trabajo "
                           "de torno a un soldador. Según lo que tiene cargado, hoy podría agarrar: "
                           + (", ".join(f.replace('_', ' ').lower() for f in familias_de[o["id"]])
                              or "ningún trabajo de máquina") + "."),
                "dato": {"id": o["id"], "nombre": o["quien"].strip(), "veces": 0,
                         "rangos": "OFICIAL", "maquinas": "", "ot_afectadas": ""},
            } for o in oficiales],
        })

    total = sum(len(g["filas"]) for g in grupos)
    if not total:
        print("No quedan trabas de máquina ni de categoría en las OT abiertas. Nada que preguntar.")
        return

    wb = Workbook()
    wb.remove(wb.active)
    _hoja_empeza(wb, total)
    _hoja_listas(wb, opciones)
    _hoja_preguntas(wb, grupos, opciones)
    _hoja_detalle(wb, grupos)
    # El orden de las solapas es el orden en que se lee: instrucciones, preguntas, resto.
    wb.move_sheet("PREGUNTAS", offset=-1)

    os.makedirs(os.path.dirname(SALIDA) or ".", exist_ok=True)
    wb.save(SALIDA)
    for g in grupos:
        print(f"  {len(g['filas']):>3}  {g['titulo']}")
    print(f"\n{total} preguntas → {SALIDA}")


if __name__ == "__main__":
    asyncio.run(main())
